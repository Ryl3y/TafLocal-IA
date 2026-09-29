"""
Lecture d'un CV (PDF, DOCX, TXT) et extraction d'informations structurées.
"""

import logging
import re
from dataclasses import dataclass, field
from datetime import date
from typing import BinaryIO

from .text import normalize

logger = logging.getLogger(__name__)


class CVParsingError(Exception):
    """Le fichier ne peut pas être lu."""


def extract_text(file_obj: BinaryIO, filename: str) -> str:
    """Extraire le texte brut d'un fichier de CV."""
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)
        if extension == "pdf":
            return _extract_pdf(file_obj)
        if extension == "docx":
            return _extract_docx(file_obj)
        if extension in {"txt", "md"}:
            return file_obj.read().decode("utf-8", errors="ignore")
        if extension == "doc":
            return _extract_legacy_doc(file_obj)
    except CVParsingError:
        raise
    except Exception as exc:  # fichier corrompu, chiffré...
        logger.warning("Lecture du CV %s impossible : %s", filename, exc)
        raise CVParsingError(f"Impossible de lire le fichier « {filename} ».") from exc
    raise CVParsingError("Format de fichier non pris en charge (PDF, DOCX, DOC ou TXT attendu).")


def _extract_pdf(file_obj: BinaryIO) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_obj)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise CVParsingError("Le PDF est protégé par un mot de passe.") from exc
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx(file_obj: BinaryIO) -> str:
    from docx import Document

    document = Document(file_obj)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def _extract_legacy_doc(file_obj: BinaryIO) -> str:
    """Les .doc (Word 97-2003) sont binaires : on récupère les suites de caractères lisibles."""
    raw = file_obj.read()
    chunks = re.findall(rb"[\x20-\x7e\xc0-\xff\n\r\t]{4,}", raw)
    text = b"\n".join(chunks).decode("latin-1", errors="ignore")
    return text.strip()


# --------------------------------------------------------------------------- #
# Sections
# --------------------------------------------------------------------------- #

SECTION_KEYWORDS = {
    "profil": ("profil", "resume", "summary", "a propos", "objectif", "about me", "presentation"),
    "experience": ("experience", "experiences", "parcours professionnel", "work experience",
                   "employment", "emplois", "historique professionnel", "stages"),
    "formation": ("formation", "formations", "education", "diplome", "diplomes", "etudes",
                  "cursus", "parcours academique", "scolarite"),
    "competences": ("competence", "competences", "skills", "savoir-faire", "aptitudes",
                    "connaissances", "outils", "technologies"),
    "langues": ("langue", "langues", "languages"),
    "projets": ("projet", "projets", "projects", "realisations"),
    "certifications": ("certification", "certifications", "certificats"),
    "interets": ("centres d'interet", "loisirs", "hobbies", "interets", "activites"),
}


def split_sections(text: str) -> dict[str, str]:
    """Découper le CV en sections selon les titres reconnus."""
    sections: dict[str, list[str]] = {"entete": []}
    current = "entete"
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        heading = normalize(stripped).strip(" :-•*#")
        detected = None
        if len(heading) <= 40:
            for section, keywords in SECTION_KEYWORDS.items():
                if any(heading == k or heading.startswith(k + " ") or heading.startswith(k + "s")
                       or heading.endswith(" " + k) for k in keywords):
                    detected = section
                    break
        if detected:
            current = detected
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(stripped)
    return {name: "\n".join(lines) for name, lines in sections.items() if lines or name != "entete"}


# --------------------------------------------------------------------------- #
# Coordonnées
# --------------------------------------------------------------------------- #

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+|00)?\d[\d\s.\-()]{7,}\d")
LINKEDIN_RE = re.compile(r"linkedin\.com/[\w\-/%]+", re.IGNORECASE)
GITHUB_RE = re.compile(r"github\.com/[\w\-]+", re.IGNORECASE)
URL_RE = re.compile(r"https?://[^\s]+", re.IGNORECASE)


_YEAR_PAIR_RE = re.compile(r"(?:19|20)\d{2}(?:19|20)\d{2}")


def _is_phone(candidate: str) -> bool:
    digits = re.sub(r"\D", "", candidate)
    # « 2022 - 2023 » : une période, pas un numéro.
    return 8 <= len(digits) <= 15 and not _YEAR_PAIR_RE.fullmatch(digits)


def extract_contacts(text: str) -> dict:
    phones = [p for p in PHONE_RE.findall(text) if _is_phone(p)]
    email = EMAIL_RE.search(text)
    linkedin = LINKEDIN_RE.search(text)
    github = GITHUB_RE.search(text)
    urls = [u for u in URL_RE.findall(text) if "linkedin" not in u and "github" not in u]
    return {
        "email": email.group(0) if email else None,
        "telephone": phones[0].strip() if phones else None,
        "linkedin": linkedin.group(0) if linkedin else None,
        "github": github.group(0) if github else None,
        "site_web": urls[0] if urls else None,
    }


# --------------------------------------------------------------------------- #
# Expérience
# --------------------------------------------------------------------------- #

MONTHS = {
    "jan": 1, "janv": 1, "janvier": 1, "january": 1, "feb": 2, "fev": 2, "fevr": 2, "fevrier": 2,
    "february": 2, "mar": 3, "mars": 3, "march": 3, "avr": 4, "avril": 4, "apr": 4, "april": 4,
    "mai": 5, "may": 5, "juin": 6, "jun": 6, "june": 6, "juil": 7, "juillet": 7, "jul": 7,
    "july": 7, "aou": 8, "aout": 8, "aug": 8, "august": 8, "sep": 9, "sept": 9, "septembre": 9,
    "september": 9, "oct": 10, "octobre": 10, "october": 10, "nov": 11, "novembre": 11,
    "november": 11, "dec": 12, "decembre": 12, "december": 12,
}
_MONTH_NAMES = "|".join(sorted(MONTHS, key=len, reverse=True))


def _date_pattern(prefix: str) -> str:
    return (
        rf"(?:(?:(?P<{prefix}m>\d{{1,2}})[/.-]|(?P<{prefix}mn>{_MONTH_NAMES})\.?\s+)?"
        rf"(?P<{prefix}y>(?:19|20)\d{{2}}))"
    )


_PRESENT = r"(?:present|aujourd'hui|aujourdhui|ce jour|actuel|actuellement|en cours|now|current|today)"
RANGE_RE = re.compile(
    _date_pattern("a")
    + r"\s*(?:-|–|—|a|au|to|jusqu'a|→)\s*(?:"
    + _date_pattern("b")
    + r"|(?P<present>" + _PRESENT + r"))"
)
EXPLICIT_YEARS_RE = re.compile(
    r"(\d{1,2})\s*\+?\s*(?:ans|annees|years?)\s+(?:d'|de\s+)?(?:experience|experiences|expertise)"
)


def _month(groups: dict, prefix: str, default: int) -> int:
    if groups.get(f"{prefix}m"):
        value = int(groups[f"{prefix}m"])
        return value if 1 <= value <= 12 else default
    if groups.get(f"{prefix}mn"):
        return MONTHS.get(groups[f"{prefix}mn"], default)
    return default


def extract_date_ranges(text: str, today: date | None = None) -> list[tuple[int, int]]:
    """Retourner les périodes (en mois absolus) trouvées dans le texte."""
    today = today or date.today()
    ranges = []
    for match in RANGE_RE.finditer(normalize(text)):
        groups = match.groupdict()
        start_year = int(groups["ay"])
        start = start_year * 12 + _month(groups, "a", 1) - 1
        if groups.get("present"):
            end = today.year * 12 + today.month - 1
        elif groups.get("by"):
            end = int(groups["by"]) * 12 + _month(groups, "b", 12) - 1
        else:
            continue
        if start <= end <= today.year * 12 + today.month and end - start <= 45 * 12:
            ranges.append((start, end))
    return ranges


def _merge_months(ranges: list[tuple[int, int]]) -> int:
    total, current_start, current_end = 0, None, None
    for start, end in sorted(ranges):
        if current_end is None or start > current_end:
            if current_end is not None:
                total += current_end - current_start + 1
            current_start, current_end = start, end
        else:
            current_end = max(current_end, end)
    if current_end is not None:
        total += current_end - current_start + 1
    return total


def estimate_experience_years(text: str, sections: dict[str, str] | None = None) -> float:
    """Estimer le nombre d'années d'expérience professionnelle."""
    normalized = normalize(text)
    explicit = [int(v) for v in EXPLICIT_YEARS_RE.findall(normalized) if 0 < int(v) <= 45]
    sections = sections or {}
    experience_text = sections.get("experience") or ""
    if not experience_text:
        # Sans section dédiée, on exclut au moins la formation.
        experience_text = "\n".join(v for k, v in sections.items() if k != "formation") or text
    months = _merge_months(extract_date_ranges(experience_text))
    from_ranges = round(months / 12, 1)
    if explicit:
        return float(max(max(explicit), from_ranges))
    return from_ranges


# --------------------------------------------------------------------------- #
# Formation
# --------------------------------------------------------------------------- #

# niveau = nombre d'années après le baccalauréat
EDUCATION_LEVELS = [
    (8, "Doctorat", ("doctorat", "phd", "ph.d", "doctorate", "bac+8", "bac + 8")),
    (5, "Bac+5 (Master / Ingénieur)", ("master", "mastere", "msc", "mba", "ingenieur", "engineer",
                                        "bac+5", "bac + 5", "dess", "dea", "m2")),
    (4, "Bac+4 (Maîtrise)", ("maitrise", "bac+4", "bac + 4", "m1")),
    (3, "Bac+3 (Licence / Bachelor)", ("licence", "bachelor", "bac+3", "bac + 3", "bsc", "l3")),
    (2, "Bac+2 (BTS / DUT / HND)", ("bts", "dut", "hnd", "deug", "bac+2", "bac + 2", "dts", "deust")),
    (0, "Baccalauréat", ("baccalaureat", "bac ", "gce a level", "a level", "probatoire")),
]


def education_level_from_text(text: str | None) -> tuple[int | None, str | None]:
    """Retourner le niveau d'études le plus élevé mentionné dans un texte."""
    normalized = f" {normalize(text)} "
    if not normalized.strip():
        return None, None
    for level, label, keywords in EDUCATION_LEVELS:
        for keyword in keywords:
            if re.search(rf"(?<![a-z0-9]){re.escape(keyword.strip())}(?![a-z0-9])", normalized):
                return level, label
    return None, None


@dataclass
class ParsedCV:
    text: str
    sections: dict[str, str] = field(default_factory=dict)
    contacts: dict = field(default_factory=dict)
    experience_years: float = 0.0
    education_level: int | None = None
    education_label: str | None = None
    word_count: int = 0


def parse_cv_text(text: str) -> ParsedCV:
    from .text import word_count

    sections = split_sections(text)
    education_source = sections.get("formation") or text
    level, label = education_level_from_text(education_source)
    if level is None and sections.get("formation"):
        level, label = education_level_from_text(text)
    return ParsedCV(
        text=text,
        sections=sections,
        contacts=extract_contacts(text),
        experience_years=estimate_experience_years(text, sections),
        education_level=level,
        education_label=label,
        word_count=word_count(text),
    )
